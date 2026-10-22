#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Módulo de Exportación y Administración de Hablantes.
Genera TXT, SRT, DOCX y JSON estructurado.
Permite renombrar hablantes en < 1 segundo sin reprocesar audio.
"""

import json
from pathlib import Path
from typing import List, Dict, Any, Optional

from src.config import formatear_tiempo, formatear_tiempo_srt


def aplicar_mapeo_hablante(speaker_raw: str, speaker_mapping: Optional[Dict[str, str]]) -> str:
    """Aplica el nombre personalizado a la etiqueta de hablante (capa de presentación)."""
    if not speaker_mapping:
        return speaker_raw
    return speaker_mapping.get(speaker_raw, speaker_raw)


def generar_txt(segments: List[Dict[str, Any]], output_path: Path, speaker_mapping: Optional[Dict[str, str]] = None) -> Path:
    """Genera archivo TXT limpio con timestamp + hablante + texto."""
    with open(output_path, "w", encoding="utf-8") as f:
        for seg in segments:
            ts = formatear_tiempo(seg["start"])
            raw_speaker = seg.get("speaker", "SPEAKER 01")
            display_speaker = aplicar_mapeo_hablante(raw_speaker, speaker_mapping)
            text = seg.get("text", "").strip()
            f.write(f"[{ts}] {display_speaker}:\n{text}\n\n")
    return output_path


def generar_srt(segments: List[Dict[str, Any]], output_path: Path, speaker_mapping: Optional[Dict[str, str]] = None) -> Path:
    """Genera archivo SRT estándar de subtítulos con timestamps exactos."""
    with open(output_path, "w", encoding="utf-8") as f:
        for i, seg in enumerate(segments, 1):
            start_ts = formatear_tiempo_srt(seg["start"])
            end_ts = formatear_tiempo_srt(seg["end"])
            raw_speaker = seg.get("speaker", "SPEAKER 01")
            display_speaker = aplicar_mapeo_hablante(raw_speaker, speaker_mapping)
            text = seg.get("text", "").strip()
            
            f.write(f"{i}\n")
            f.write(f"{start_ts} --> {end_ts}\n")
            f.write(f"[{display_speaker}] {text}\n\n")
    return output_path


def generar_docx(segments: List[Dict[str, Any]], output_path: Path, titulo: str = "Transcripción", speaker_mapping: Optional[Dict[str, str]] = None) -> Path:
    """Genera documento DOCX profesional con estilos legibles."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
    except ImportError:
        return output_path

    doc = Document()
    
    # Configurar márgenes limpios
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Título principal
    h = doc.add_heading(titulo, level=1)
    
    current_speaker = None
    first_ts = "00:00:00"
    accumulated_texts = []

    def flush_paragraph():
        nonlocal current_speaker, first_ts, accumulated_texts
        if current_speaker is not None and accumulated_texts:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            
            # Nombre de hablante + timestamp en negrita y color azul institucional
            display_speaker = aplicar_mapeo_hablante(current_speaker, speaker_mapping)
            run_header = p.add_run(f"[{first_ts}] {display_speaker}:\n")
            run_header.bold = True
            run_header.font.size = Pt(11)
            run_header.font.color.rgb = RGBColor(0, 51, 102)
            
            # Contenido de la intervención
            run_body = p.add_run(" ".join(accumulated_texts))
            run_body.font.size = Pt(11)
            run_body.font.color.rgb = RGBColor(30, 30, 30)

    for seg in segments:
        speaker = seg.get("speaker", "SPEAKER 01")
        ts = formatear_tiempo(seg["start"])
        text = seg.get("text", "").strip()
        
        if speaker != current_speaker:
            flush_paragraph()
            current_speaker = speaker
            first_ts = ts
            accumulated_texts = [text]
        else:
            accumulated_texts.append(text)

    flush_paragraph()
    doc.save(str(output_path))
    return output_path


def generar_json(segments: List[Dict[str, Any]], output_path: Path, metadata: Optional[Dict[str, Any]] = None, speaker_mapping: Optional[Dict[str, str]] = None) -> Path:
    """Genera archivo JSON estructurado con toda la metadata para visualización y edición."""
    data = {
        "metadata": metadata or {},
        "speaker_mapping": speaker_mapping or {},
        "segments": segments,
    }
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return output_path


def exportar_todo(segments: List[Dict[str, Any]], output_dir: Path, base_name: str, speaker_mapping: Optional[Dict[str, str]] = None, metadata: Optional[Dict[str, Any]] = None) -> Dict[str, Path]:
    """Genera simultáneamente todos los formatos (TXT, SRT, DOCX, JSON)."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    
    txt_path = output_dir / f"{base_name}.txt"
    srt_path = output_dir / f"{base_name}.srt"
    docx_path = output_dir / f"{base_name}.docx"
    json_path = output_dir / f"{base_name}_completo.json"
    
    generar_txt(segments, txt_path, speaker_mapping)
    generar_srt(segments, srt_path, speaker_mapping)
    generar_docx(segments, docx_path, titulo=f"Transcripción: {base_name}", speaker_mapping=speaker_mapping)
    generar_json(segments, json_path, metadata=metadata, speaker_mapping=speaker_mapping)
    
    return {
        "txt": txt_path,
        "srt": srt_path,
        "docx": docx_path,
        "json": json_path,
    }


def reexportar_desde_json(json_path: Path, nuevo_mapping: Dict[str, str], output_dir: Optional[Path] = None) -> Dict[str, Path]:
    """
    Carga un archivo JSON de transcripción previa, actualiza los nombres de los hablantes
    y reexporta TXT, SRT, DOCX y JSON en menos de 1 segundo sin tocar Whisper ni pyannote.
    """
    json_path = Path(json_path)
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    
    segments = data.get("segments") or data.get("segmentos") or []
    metadata = data.get("metadata", {})
    base_name = json_path.stem.replace("_completo", "")
    target_dir = Path(output_dir) if output_dir else json_path.parent
    
    return exportar_todo(
        segments=segments,
        output_dir=target_dir,
        base_name=base_name,
        speaker_mapping=nuevo_mapping,
        metadata=metadata,
    )
